#!/usr/bin/env python3
"""
Recreate all PowerPoint presentations with proper 16:9 widescreen format
"""

import os
from pathlib import Path
from bs4 import BeautifulSoup
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def extract_text_from_element(element):
    """Extract clean text from HTML element, preserving structure"""
    text = element.get_text(separator='\n', strip=True)
    lines = [line.strip() for line in text.split('\n') if line.strip()]
    return '\n'.join(lines)

def parse_html_slides(html_file):
    """Parse HTML file and extract slide content"""
    with open(html_file, 'r', encoding='utf-8') as f:
        soup = BeautifulSoup(f.read(), 'html.parser')

    slides_data = []
    slides = soup.find_all('div', class_='slide')

    for slide in slides:
        slide_info = {
            'type': 'content',
            'title': '',
            'content': [],
            'is_title_slide': False,
            'is_practice_slide': False
        }

        # Check for h1 (main title slides)
        h1 = slide.find('h1')
        if h1:
            slide_info['title'] = h1.get_text(strip=True)
            slide_info['is_title_slide'] = True

            h2 = slide.find('h2')
            if h2:
                slide_info['content'].append(h2.get_text(strip=True))

            paragraphs = slide.find_all('p')
            for p in paragraphs:
                text = p.get_text(strip=True)
                if text and text not in slide_info['content']:
                    slide_info['content'].append(text)
        else:
            # Regular content slide
            h2 = slide.find('h2')
            if h2:
                slide_info['title'] = h2.get_text(strip=True)
                slide_info['title'] = ''.join(char for char in slide_info['title']
                                              if ord(char) < 0x1F600 or ord(char) > 0x1F6FF)
                slide_info['title'] = slide_info['title'].strip()

            for div in slide.find_all('div', recursive=True):
                if 'slide' in div.get('class', []):
                    continue

                h3 = div.find('h3', recursive=False)
                if h3:
                    header_text = h3.get_text(strip=True)
                    header_text = ''.join(char for char in header_text
                                         if ord(char) < 0x1F600 or ord(char) > 0x1F6FF)
                    if header_text:
                        slide_info['content'].append(header_text.strip())

                paragraphs = div.find_all('p', recursive=False)
                for p in paragraphs:
                    text = p.get_text(strip=True)
                    text = ''.join(char for char in text
                                  if ord(char) < 0x1F600 or ord(char) > 0x1F6FF)
                    if text and len(text) > 3:
                        slide_info['content'].append(text.strip())

            lists = slide.find_all(['ul', 'ol'])
            for ul in lists:
                items = ul.find_all('li')
                for item in items:
                    text = item.get_text(strip=True)
                    text = ''.join(char for char in text
                                  if ord(char) < 0x1F600 or ord(char) > 0x1F6FF)
                    if text and len(text) > 3:
                        slide_info['content'].append('• ' + text.strip())

        if 'practice' in slide_info['title'].lower() or 'time to practice' in slide_info['title'].lower():
            slide_info['is_practice_slide'] = True

        slides_data.append(slide_info)

    return slides_data

def add_title_slide(prs, slide_info):
    """Add a title slide to the presentation"""
    slide_layout = prs.slide_layouts[6]  # Blank layout
    slide = prs.slides.add_slide(slide_layout)

    # Add background
    left = top = Inches(0)
    width = prs.slide_width
    height = prs.slide_height

    bg_shape = slide.shapes.add_shape(1, left, top, width, height)
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = RGBColor(240, 248, 255)  # Light blue
    bg_shape.line.fill.background()

    # Add title - centered for widescreen
    title_box = slide.shapes.add_textbox(
        Inches(1), Inches(2), Inches(11.33), Inches(1.5)
    )
    title_frame = title_box.text_frame
    title_frame.text = slide_info['title']
    title_para = title_frame.paragraphs[0]
    title_para.font.size = Pt(54)
    title_para.font.bold = True
    title_para.font.color.rgb = RGBColor(0, 51, 102)
    title_para.alignment = PP_ALIGN.CENTER

    # Add subtitle/content
    if slide_info['content']:
        y_pos = 4.0
        for content in slide_info['content'][:2]:
            subtitle_box = slide.shapes.add_textbox(
                Inches(1), Inches(y_pos), Inches(11.33), Inches(0.6)
            )
            subtitle_frame = subtitle_box.text_frame
            subtitle_frame.text = content
            subtitle_para = subtitle_frame.paragraphs[0]
            subtitle_para.font.size = Pt(28)
            subtitle_para.font.color.rgb = RGBColor(70, 130, 180)
            subtitle_para.alignment = PP_ALIGN.CENTER
            y_pos += 0.7

def add_content_slide(prs, slide_info):
    """Add a content slide with title and bullet points"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    # Add white background
    left = top = Inches(0)
    width = prs.slide_width
    height = prs.slide_height

    bg_shape = slide.shapes.add_shape(1, left, top, width, height)
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    bg_shape.line.fill.background()

    # Add blue header bar
    header_shape = slide.shapes.add_shape(
        1, Inches(0), Inches(0), prs.slide_width, Inches(0.15)
    )
    header_shape.fill.solid()
    header_shape.fill.fore_color.rgb = RGBColor(0, 102, 204)
    header_shape.line.fill.background()

    # Add title
    title_box = slide.shapes.add_textbox(
        Inches(0.5), Inches(0.4), Inches(12), Inches(0.8)
    )
    title_frame = title_box.text_frame
    title_frame.text = slide_info['title']
    title_para = title_frame.paragraphs[0]
    title_para.font.size = Pt(40)
    title_para.font.bold = True
    title_para.font.color.rgb = RGBColor(0, 51, 102)

    # Add blue underline
    underline_shape = slide.shapes.add_shape(
        1, Inches(0.5), Inches(1.1), Inches(2), Inches(0.05)
    )
    underline_shape.fill.solid()
    underline_shape.fill.fore_color.rgb = RGBColor(0, 102, 204)
    underline_shape.line.fill.background()

    # Add content
    if slide_info['content']:
        content_box = slide.shapes.add_textbox(
            Inches(0.7), Inches(1.5), Inches(11.8), Inches(5)
        )
        content_frame = content_box.text_frame
        content_frame.word_wrap = True

        for i, content in enumerate(slide_info['content'][:10]):
            if i == 0:
                p = content_frame.paragraphs[0]
            else:
                p = content_frame.add_paragraph()

            p.text = content
            p.font.size = Pt(18)
            p.font.color.rgb = RGBColor(51, 51, 51)
            p.space_after = Pt(8)

            if content.startswith('•'):
                p.level = 1
            else:
                p.level = 0
                p.font.bold = True
                p.font.size = Pt(20)

def add_practice_slide(prs, slide_info):
    """Add a practice/end slide with blue background"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    # Add blue background
    left = top = Inches(0)
    width = prs.slide_width
    height = prs.slide_height

    bg_shape = slide.shapes.add_shape(1, left, top, width, height)
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = RGBColor(0, 102, 204)
    bg_shape.line.fill.background()

    # Add title - centered for widescreen
    title_box = slide.shapes.add_textbox(
        Inches(1), Inches(2.5), Inches(11.33), Inches(1.5)
    )
    title_frame = title_box.text_frame
    title_frame.text = slide_info['title']
    title_para = title_frame.paragraphs[0]
    title_para.font.size = Pt(48)
    title_para.font.bold = True
    title_para.font.color.rgb = RGBColor(255, 255, 255)
    title_para.alignment = PP_ALIGN.CENTER

    # Add content
    if slide_info['content']:
        content_box = slide.shapes.add_textbox(
            Inches(1), Inches(4.2), Inches(11.33), Inches(1)
        )
        content_frame = content_box.text_frame
        content_frame.text = slide_info['content'][0] if slide_info['content'] else ""
        content_para = content_frame.paragraphs[0]
        content_para.font.size = Pt(24)
        content_para.font.color.rgb = RGBColor(255, 255, 255)
        content_para.alignment = PP_ALIGN.CENTER

def create_pptx_from_html(html_file, output_file):
    """Create PowerPoint presentation from HTML file with 16:9 format"""
    print(f"Processing {html_file}...")

    slides_data = parse_html_slides(html_file)
    print(f"  Found {len(slides_data)} slides")

    # Create presentation with 16:9 widescreen format
    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 16:9 widescreen
    prs.slide_height = Inches(7.5)

    # Add slides
    for i, slide_info in enumerate(slides_data):
        print(f"  Creating slide {i+1}: {slide_info['title'][:50]}...")

        if slide_info['is_title_slide']:
            add_title_slide(prs, slide_info)
        elif slide_info['is_practice_slide']:
            add_practice_slide(prs, slide_info)
        else:
            add_content_slide(prs, slide_info)

    prs.save(output_file)
    print(f"  Saved to {output_file}\n")

def main():
    """Recreate all PowerPoint presentations with proper 16:9 format"""
    chapters = [
        ('chapter_01_introduction', 'Introduction to Python'),
        ('chapter_02_working_with_data', 'Working with Data'),
        ('chapter_03_making_decisions', 'Making Decisions'),
        ('chapter_04_loops', 'Loops'),
        ('chapter_05_lists', 'Lists'),
        ('chapter_06_dictionaries', 'Dictionaries'),
        ('chapter_07_functions', 'Functions'),
        ('chapter_08_tuples_and_sets', 'Tuples and Sets'),
        ('chapter_09_file_io', 'File I/O'),
        ('chapter_10_error_handling', 'Error Handling')
    ]

    base_dir = Path('/home/user/courses/python_course')

    print("Recreating ALL PowerPoint presentations with 16:9 widescreen format\n")
    print("=" * 70)

    for chapter_dir, chapter_name in chapters:
        chapter_path = base_dir / chapter_dir
        html_file = chapter_path / 'presentation.html'
        pptx_file = chapter_path / 'presentation.pptx'

        if not html_file.exists():
            print(f"Warning: {html_file} not found, skipping...")
            continue

        create_pptx_from_html(str(html_file), str(pptx_file))

    print("=" * 70)
    print("All PowerPoint presentations recreated with 16:9 widescreen format!")
    print(f"New dimensions: 13.333\" x 7.5\" (16:9 ratio)")

if __name__ == '__main__':
    main()
