#!/usr/bin/env python3
"""
HTML to PowerPoint Converter
Converts HTML slide presentations to PPTX format
"""

import os
import sys
from pathlib import Path

def install_requirements():
    """Install required packages"""
    import subprocess

    packages = [
        'python-pptx',
        'beautifulsoup4',
        'pillow',
        'selenium',
        'webdriver-manager'
    ]

    print("Installing required packages...")
    for package in packages:
        subprocess.run([sys.executable, '-m', 'pip', 'install', '-q', package])
    print("✓ Packages installed!")

def html_to_pptx_screenshot(html_file, output_file):
    """
    Convert HTML to PPTX by taking screenshots of each slide
    Requires Chrome/Chromium browser
    """
    from selenium import webdriver
    from selenium.webdriver.chrome.service import Service
    from selenium.webdriver.chrome.options import Options
    from webdriver_manager.chrome import ChromeDriverManager
    from pptx import Presentation
    from pptx.util import Inches
    from PIL import Image
    import io

    # Setup Chrome options
    chrome_options = Options()
    chrome_options.add_argument('--headless')
    chrome_options.add_argument('--no-sandbox')
    chrome_options.add_argument('--disable-dev-shm-usage')
    chrome_options.add_argument('--window-size=1280,720')

    print(f"Converting {html_file} to {output_file}...")

    # Initialize Chrome driver
    driver = webdriver.Chrome(
        service=Service(ChromeDriverManager().install()),
        options=chrome_options
    )

    try:
        # Load HTML file
        file_url = f'file://{os.path.abspath(html_file)}'
        driver.get(file_url)

        # Find all slides
        from selenium.webdriver.common.by import By
        slides = driver.find_elements(By.CLASS_NAME, 'slide')

        if not slides:
            print("❌ No slides found in HTML file")
            return False

        print(f"Found {len(slides)} slides")

        # Create presentation
        prs = Presentation()
        prs.slide_width = Inches(10)  # 1280px
        prs.slide_height = Inches(5.625)  # 720px

        # Screenshot each slide
        for i, slide in enumerate(slides, 1):
            print(f"  Processing slide {i}/{len(slides)}...")

            # Scroll to slide
            driver.execute_script("arguments[0].scrollIntoView(true);", slide)

            # Take screenshot
            screenshot = slide.screenshot_as_png

            # Add to presentation
            blank_slide_layout = prs.slide_layouts[6]  # Blank layout
            pptx_slide = prs.slides.add_slide(blank_slide_layout)

            # Add image to slide
            img_stream = io.BytesIO(screenshot)
            left = top = Inches(0)
            pptx_slide.shapes.add_picture(img_stream, left, top,
                                         width=prs.slide_width,
                                         height=prs.slide_height)

        # Save presentation
        prs.save(output_file)
        print(f"✓ Saved to {output_file}")
        return True

    except Exception as e:
        print(f"❌ Error: {e}")
        return False
    finally:
        driver.quit()

def convert_all_presentations(course_dir):
    """Convert all HTML presentations in the course directory"""
    course_path = Path(course_dir)

    # Find all presentation.html files
    html_files = list(course_path.rglob('presentation.html'))

    if not html_files:
        print("No presentation.html files found!")
        return

    print(f"\nFound {len(html_files)} presentations to convert\n")

    for html_file in html_files:
        # Create output filename
        output_file = html_file.parent / 'presentation.pptx'

        # Convert
        html_to_pptx_screenshot(str(html_file), str(output_file))
        print()

def main():
    # Check if requirements are installed
    try:
        import pptx
        from bs4 import BeautifulSoup
        from selenium import webdriver
    except ImportError:
        print("Installing required packages...")
        install_requirements()

    # Get course directory
    if len(sys.argv) > 1:
        course_dir = sys.argv[1]
    else:
        # Default to python_course directory
        course_dir = os.path.join(os.path.dirname(__file__), '..')

    convert_all_presentations(course_dir)
    print("\n✓ All conversions complete!")

if __name__ == '__main__':
    main()
